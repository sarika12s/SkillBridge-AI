"""Comprehensive automated test suite for resume upload, parsing, section detection, and security."""

import io
import fitz
import docx
import pytest
from fastapi.testclient import TestClient

from app.core.security import create_access_token
from app.ai.segmentation.section_classifier import is_potential_heading, segment_resume_sections
from app.ai.segmentation.entity_extractor import extract_contact_information
from app.ai.ingestion.text_quality import evaluate_text_quality, TextQualityStatus


def create_sample_pdf(content: str) -> bytes:
    """Create an in-memory valid text-based PDF using PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(fitz.Point(40, 50), content)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx(content_lines: list[tuple[str, int]]) -> bytes:
    """Create an in-memory valid DOCX using python-docx."""
    doc = docx.Document()
    for text, level in content_lines:
        if level == 0:
            doc.add_paragraph(text)
        else:
            doc.add_heading(text, level=level)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def auth_headers(client: TestClient) -> dict:
    """Register and authenticate a test user, returning Bearer auth headers."""
    register_payload = {
        "email": "candidate@college.edu",
        "password": "Password123!",
        "full_name": "Jordan Smith",
    }
    resp = client.post("/api/v1/auth/register", json=register_payload)
    if resp.status_code == 201:
        token = resp.json()["access_token"]
    else:
        login_resp = client.post(
            "/api/v1/auth/login",
            json={"email": register_payload["email"], "password": register_payload["password"]},
        )
        token = login_resp.json()["access_token"]

    return {"Authorization": f"Bearer {token}"}


# ------------------------------------------------------------------------------
# Unit Tests for Section Detection & Entity Extraction
# ------------------------------------------------------------------------------

def test_section_heading_detection_variations():
    """Verify that multiple heading variations are correctly normalized."""
    test_cases = [
        ("PROFESSIONAL SUMMARY", "SUMMARY"),
        ("Career Objective", "SUMMARY"),
        ("About Me", "SUMMARY"),
        ("ACADEMIC QUALIFICATIONS", "EDUCATION"),
        ("Educational Background", "EDUCATION"),
        ("WORK EXPERIENCE", "EXPERIENCE"),
        ("Employment History", "EXPERIENCE"),
        ("TECHNICAL SKILLS", "SKILLS"),
        ("Core Competencies", "SKILLS"),
        ("ACADEMIC PROJECTS", "PROJECTS"),
        ("LICENSES & CERTIFICATIONS", "CERTIFICATIONS"),
        ("HONORS AND AWARDS", "ACHIEVEMENTS"),
        ("LANGUAGES", "LANGUAGES"),
        ("INTERESTS", "INTERESTS"),
    ]
    for heading, expected_type in test_cases:
        is_hdr, norm_type, conf = is_potential_heading(heading)
        assert is_hdr is True, f"Failed to recognize heading: {heading}"
        assert norm_type == expected_type, f"Expected {expected_type} for {heading}, got {norm_type}"
        assert conf >= 0.85


def test_contact_information_extraction():
    """Verify extraction of name, email, phone, linkedin, and github."""
    sample_text = (
        "Alex Carter\n"
        "alex.carter@gmail.com | +1 (555) 234-5678 | San Francisco, CA\n"
        "https://linkedin.com/in/alexcarter | https://github.com/alexcarter\n"
        "https://alexcarter.dev\n\n"
        "PROFESSIONAL SUMMARY\n"
        "Skilled software engineer."
    )
    contact = extract_contact_information(sample_text)
    assert contact["name"] == "Alex Carter"
    assert contact["email"] == "alex.carter@gmail.com"
    assert contact["phone"] is not None and "555" in contact["phone"]
    assert contact["linkedin_url"] == "https://linkedin.com/in/alexcarter"
    assert contact["github_url"] == "https://github.com/alexcarter"
    assert contact["portfolio_url"] == "https://alexcarter.dev"


def test_text_quality_decision():
    """Verify deterministic quality scoring decisions."""
    # Sufficient text
    good_text = "This is a full resume document. " * 30
    status, metrics = evaluate_text_quality(good_text, is_pdf=True)
    assert status == TextQualityStatus.GOOD_TEXT

    # Empty / Scanned text
    scanned_text = ""
    status, metrics = evaluate_text_quality(scanned_text, is_pdf=True)
    assert status == TextQualityStatus.OCR_REQUIRED

    # Degraded short text (below 100 chars or 20 words)
    short_text = "Page 1 - Conf"
    status, metrics = evaluate_text_quality(short_text, is_pdf=True)
    assert status == TextQualityStatus.OCR_REQUIRED


# ------------------------------------------------------------------------------
# Integration API Tests for Resume Upload & Lifecycle
# ------------------------------------------------------------------------------

def test_valid_pdf_upload_and_parse(client: TestClient, auth_headers: dict):
    """Test full upload, section segmentation, entity extraction, and response for a PDF."""
    pdf_content = (
        "Jordan Smith\n"
        "jordan.smith@example.com | 555-432-1098 | Austin, TX\n"
        "https://linkedin.com/in/jordansmith | https://github.com/jordansmith\n\n"
        "PROFESSIONAL SUMMARY\n"
        "Experienced Backend Developer specializing in scalable cloud microservices, REST APIs, and event-driven systems.\n\n"
        "EDUCATION\n"
        "B.S. in Computer Engineering\n"
        "University of Texas at Austin - May 2021\n\n"
        "WORK EXPERIENCE\n"
        "Software Engineer at Dell Technologies\n"
        "Jun 2021 - Present\n"
        "Built distributed telemetry pipelines processing 10M records daily using FastAPI and PostgreSQL.\n\n"
        "PROJECTS\n"
        "Cloud Storage Gateway (Python, Docker, AWS S3)\n"
        "Engineered an S3 proxy cache reducing egress latency by 45%.\n\n"
        "TECHNICAL SKILLS\n"
        "Python, FastAPI, PostgreSQL, Docker, Kubernetes, Linux, Redis\n\n"
        "CERTIFICATIONS\n"
        "AWS Certified Developer Associate by Amazon Web Services\n"
    )
    pdf_bytes = create_sample_pdf(pdf_content)

    response = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("jordan_resume.pdf", pdf_bytes, "application/pdf")},
        data={"title": "Jordan Smith SWE Resume"},
        headers=auth_headers,
    )
    assert response.status_code == 201
    data = response.json()

    assert data["title"] == "Jordan Smith SWE Resume"
    assert data["file_type"] == "pdf"
    assert data["parsing_status"] == "COMPLETED"
    assert data["extraction_method"] == "TEXT"
    assert data["character_count"] > 100
    assert data["page_count"] >= 1

    # Verify extracted contact details
    personal_info = data["personal_information"]
    assert personal_info["name"] == "Jordan Smith"
    assert personal_info["email"] == "jordan.smith@example.com"
    assert personal_info["linkedin_url"] == "https://linkedin.com/in/jordansmith"
    assert personal_info["github_url"] == "https://github.com/jordansmith"

    # Verify structured sections
    section_types = [s["section_type"] for s in data["sections"]]
    assert "SUMMARY" in section_types
    assert "EDUCATION" in section_types
    assert "EXPERIENCE" in section_types
    assert "PROJECTS" in section_types
    assert "SKILLS" in section_types

    # Verify sub-entities
    assert len(data["experience"]) >= 1
    assert "Dell" in data["experience"][0]["company_name"]
    assert len(data["projects"]) >= 1
    assert len(data["certifications"]) >= 1

    resume_id = data["id"]

    # Verify GET by ID
    get_resp = client.get(f"/api/v1/resumes/{resume_id}", headers=auth_headers)
    assert get_resp.status_code == 200
    assert get_resp.json()["id"] == resume_id

    # Verify listing
    list_resp = client.get("/api/v1/resumes", headers=auth_headers)
    assert list_resp.status_code == 200
    assert any(r["id"] == resume_id for r in list_resp.json())

    # Verify DELETE
    del_resp = client.delete(f"/api/v1/resumes/{resume_id}", headers=auth_headers)
    assert del_resp.status_code == 204

    # Verify 404 after deletion
    after_del_resp = client.get(f"/api/v1/resumes/{resume_id}", headers=auth_headers)
    assert after_del_resp.status_code == 404


def test_valid_docx_upload_and_parse(client: TestClient, auth_headers: dict):
    """Test upload and parsing of a valid DOCX resume."""
    docx_lines = [
        ("Jane Doe", 1),
        ("jane.doe@columbia.edu | 212-555-0199 | New York, NY | https://github.com/janedoe", 0),
        ("CAREER OBJECTIVE", 1),
        ("Dedicated software engineer with a strong foundation in modern web frameworks and database optimization.", 0),
        ("EDUCATION", 1),
        ("B.S. in Computer Science - Columbia University, 2023", 0),
        ("WORK EXPERIENCE", 1),
        ("Frontend Developer at Tech Corp\nMay 2023 - Present\nImplemented React components with TypeScript and Tailwind CSS.", 0),
        ("TECHNICAL SKILLS", 1),
        ("TypeScript, React, Node.js, Python, PostgreSQL, Git", 0),
    ]
    docx_bytes = create_sample_docx(docx_lines)

    response = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("jane_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=auth_headers,
    )
    assert response.status_code == 201
    data = response.json()
    assert data["file_type"] == "docx"
    assert data["parsing_status"] == "COMPLETED"
    assert data["personal_information"]["email"] == "jane.doe@columbia.edu"
    assert any(s["section_type"] == "SKILLS" for s in data["sections"])


def test_upload_invalid_file_extension(client: TestClient, auth_headers: dict):
    """Verify that non-supported file formats (e.g. .txt, .exe) are rejected."""
    fake_content = b"Some plain text resume content."
    response = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("resume.txt", fake_content, "text/plain")},
        headers=auth_headers,
    )
    assert response.status_code == 422
    assert "Unsupported file format" in response.json()["detail"]


def test_upload_empty_document(client: TestClient, auth_headers: dict):
    """Verify that 0-byte files are rejected."""
    response = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("empty.pdf", b"", "application/pdf")},
        headers=auth_headers,
    )
    assert response.status_code == 422
    assert "empty" in response.json()["detail"].lower()


def test_upload_corrupted_pdf_header(client: TestClient, auth_headers: dict):
    """Verify that files claiming to be PDF without %PDF magic bytes are rejected."""
    corrupted_bytes = b"NOT_A_PDF_MAGIC_BYTES_123456789"
    response = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("corrupt.pdf", corrupted_bytes, "application/pdf")},
        headers=auth_headers,
    )
    assert response.status_code == 422
    assert "magic header" in response.json()["detail"].lower()


def test_unauthorized_upload_attempt(client: TestClient):
    """Verify that requests without a Bearer token are rejected with 401."""
    pdf_bytes = create_sample_pdf("Sample text resume content.")
    response = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("test.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 401
